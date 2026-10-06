"""Required inspection evidence per product type.

Source: "BRT_required_evidence.pdf" (Required Evidence matrix) plus the
decisions agreed with the business (2026-10):
  - PE RSV follows the RSV rule; Spark test for every RSV >= DN200
    (FL / SO / SP / PE); DN375 has no DAQ.
  - Pretap connectors follow the DI fitting rule; pretap bushes need no
    pressure test.
  - L-Type hydrant covers / heads: checklist only (for now).
  - Comp covers (Jinmeng) are not configured yet.

A product is classified from the product reference table (imported from
reference.xlsx: sheets "ReferenceData" and "Product Codes"), falling back
to item-code prefixes and description keywords.
"""
import io
import re


def _c(zh, en):
    return {'zh': zh, 'en': en}


# ── Check items (one line of the PDF each) ───────────────────────────────────
CHECK_CHECKLIST = _c('检验清单已由检验员签字并注明日期，所有项目合格',
                     'Inspection checklist signed and dated by the inspector; all items acceptable')
CHECK_MATERIAL = _c('材质报告：化学成分和力学性能在限值内',
                    'Material report: chemical and mechanical properties within limits')
CHECK_BRT_OK = _c('BRT 文件齐全且合格', 'BRT documents complete and acceptable')
CHECK_GAL_BOLTS = _c('GAL 型号：螺栓为钢制（Q235B）', 'GAL variant: bolts are steel (Q235B)')

RSV_BRT = [
    _c('材质表：化学成分和力学性能在限值内（500-7）',
       'Material sheet: chemical and mechanical within limits (500-7)'),
    _c('检验报告：C1–C7 和 A1–A9 全部 OK（参照 SPEC 表）',
       'Checking Report: C1–C7 & A1–A9 OK (refer to SPEC sheet)'),
    _c('检验报告 DAQ 合格（AB–BH 单元格）：T1 [AL]、T2 [AW] ≥ 1.76 MPa，T3 [BH] ≥ 2.4 MPa',
       'Checking Report DAQ acceptable (cells AB–BH): T1 [AL], T2 [AW] ≥ 1.76 MPa, T3 [BH] ≥ 2.4 MPa'),
]
RSV375_BRT = RSV_BRT[:2] + [
    _c('DN375 无 DAQ 数据：通过 Z–AG 单元格和 V-Trust 视频核对压力测试',
       'DN375 has no DAQ: verify the pressure test via cells Z–AG and the V-Trust video'),
]
DI_BRT = [
    _c('DPL 管件铸件检验报告：所有项目合格', 'DPL Fitting casting inspection report: all criteria acceptable'),
    _c('DPL 管件最终检验报告：所有项目合格', 'DPL Fitting final inspection report: all criteria acceptable'),
]
UMC_BRT = [
    _c('316 SS.jpg：确认材质为 316 不锈钢', '316 SS.jpg: confirm 316 stainless steel'),
    _c('装配报告：所有项目合格', 'Assembly report: all criteria acceptable'),
    _c('Bolt 316.jpg：确认螺栓为 316 不锈钢', 'Bolt 316.jpg: confirm bolts are 316'),
    _c('尺寸报告：与 Daemco 设计图纸比对（UMC - 2019-00-00\\Design）',
       'Dimension report: compare with the Daemco design drawings (UMC - 2019-00-00\\Design)'),
]
CLAMP_BRT = [
    _c('装配文件夹：所有勾选项合格', 'Assembly folder files: all check boxes acceptable'),
    _c('尺寸报告：与 REPAIR CLAMP 2023.11.7 (Bolt Info).xlsx 比对',
       'Dimension report: compare with REPAIR CLAMP 2023.11.7 (Bolt Info).xlsx'),
    _c('材质文件夹：确认为 316 不锈钢', 'Material folder files: check for 316'),
]
PE_GRIPPA_BRT = [
    _c('装配报告：所有勾选项合格', 'Assembly report: all check boxes acceptable'),
    _c('涂层报告：所有勾选项合格', 'Coating report: all check boxes acceptable'),
    _c('尺寸检验报告：测量值与图纸一致', 'Dimensional inspection report: measurements match the print'),
    _c('材质报告：化学成分和力学性能在限值内', 'Material report: chemical and mechanical within limits'),
]

# DAQ thresholds (MPa) for the automatic check on RSV pressure tests
DAQ_LIMITS = (('t1', _c('T1 平均值 [AL] 闸板测试 1', 'T1 average [AL] gate test 1'), 1.76),
              ('t2', _c('T2 平均值 [AW] 闸板测试 2', 'T2 average [AW] gate test 2'), 1.76),
              ('t3', _c('T3 平均值 [BH] 阀体测试', 'T3 average [BH] body test'), 2.4))


def _ev(etype, *checks):
    return {'type': etype, 'checks': list(checks)}


# ── Product types (rows of the PDF) ──────────────────────────────────────────
PRODUCT_TYPES = {
    'rsv_small': {'name': _c('闸阀 ≤ DN180（FL/SO/SP/PE）', 'RSV ≤ DN180 (FL/SO/SP/PE)'),
                  'evidence': [_ev('brt', *RSV_BRT), _ev('daq'),
                               _ev('vtrust', _c('V-Trust 压力测试视频齐全，结果合格',
                                                'V-Trust pressure test videos received and acceptable'))]},
    'rsv_large': {'name': _c('闸阀 DN200–DN300（FL/SO/SP/PE）', 'RSV DN200–DN300 (FL/SO/SP/PE)'),
                  'evidence': [_ev('brt', *RSV_BRT),
                               _ev('spark', _c('电火花测试视频全部收到', 'All spark test videos received')),
                               _ev('daq'),
                               _ev('vtrust', _c('V-Trust 压力测试视频齐全，结果合格',
                                                'V-Trust pressure test videos received and acceptable'))]},
    'rsv_375': {'name': _c('闸阀 DN375', 'RSV DN375'),
                'evidence': [_ev('brt', *RSV375_BRT),
                             _ev('spark', _c('电火花测试视频全部收到', 'All spark test videos received')),
                             _ev('vtrust', _c('V-Trust 压力测试视频齐全，结果合格',
                                              'V-Trust pressure test videos received and acceptable'))]},
    'valve_cap': {'name': _c('阀帽', 'Valve caps'), 'evidence': [_ev('checklist', CHECK_CHECKLIST)]},
    'umc': {'name': _c('UMC 联轴器', 'UMC coupling'), 'evidence': [_ev('brt', *UMC_BRT)]},
    'umc_gal': {'name': _c('UMC GAL 联轴器', 'UMC GAL coupling'),
                'evidence': [_ev('brt', *UMC_BRT, CHECK_GAL_BOLTS)]},
    'gibault': {'name': _c('Gibault 联轴器', 'Gibault joint coupling'),
                'evidence': [_ev('brt', CHECK_BRT_OK), _ev('material', CHECK_MATERIAL)]},
    'di_fitting': {'name': _c('球铁管件', 'DI fitting'),
                   'evidence': [_ev('brt', *DI_BRT),
                                _ev('pressure', _c('Reliable 压力测试记录：保压 20 秒',
                                                   'Pressure test records (at Reliable): 20-second pressure holding'))]},
    'pretap_bush': {'name': _c('预开孔衬套', 'Pretap bush'), 'evidence': [_ev('brt', *DI_BRT)]},
    'blank_flange': {'name': _c('盲板 / 开孔盲板', 'Blank / tapped flange'),
                     'evidence': [_ev('brt', CHECK_BRT_OK), _ev('material', CHECK_MATERIAL)]},
    'pe_grippa': {'name': _c('PE Grippa', 'PE Grippa'), 'evidence': [_ev('brt', *PE_GRIPPA_BRT)]},
    'repair_clamp': {'name': _c('修补卡箍（Repair/Boss/Mini/Circle）', 'Repair / Boss / Mini / Circle clamp'),
                     'evidence': [_ev('brt', *CLAMP_BRT),
                                  _ev('xrf', _c('XRF 报告：确认为 316 不锈钢', 'XRF report: check for 316'))]},
    'repair_clamp_gal': {'name': _c('修补卡箍 GAL', 'Repair clamp GAL'),
                         'evidence': [_ev('brt', *CLAMP_BRT, CHECK_GAL_BOLTS),
                                      _ev('xrf', _c('XRF 报告：确认为 316 不锈钢', 'XRF report: check for 316'))]},
    'spindle': {'name': _c('加长杆', 'Extension spindle'),
                'evidence': [_ev('checklist', CHECK_CHECKLIST), _ev('material', CHECK_MATERIAL)]},
    'ss_strap': {'name': _c('不锈钢带', 'Stainless steel strap'),
                 'evidence': [_ev('checklist', CHECK_CHECKLIST), _ev('material', CHECK_MATERIAL)]},
    'gasket': {'name': _c('垫片', 'Gasket'), 'evidence': [_ev('checklist', CHECK_CHECKLIST)]},
    'handwheel': {'name': _c('手轮', 'Handwheel'), 'evidence': [_ev('checklist', CHECK_CHECKLIST)]},
    'cover': {'name': _c('井盖 / 阀箱 / 盖板', 'Cover / box / lid'),
              'evidence': [_ev('checklist', CHECK_CHECKLIST), _ev('material', CHECK_MATERIAL)]},
    'l_type_head': {'name': _c('L 型消防栓头', 'L-Type hydrant head'),
                    'evidence': [_ev('checklist', CHECK_CHECKLIST)]},
    'l_type': {'name': _c('L 型消防栓盖', 'L-Type hydrant cover'),
               'evidence': [_ev('checklist', CHECK_CHECKLIST)]},
    'latch_pin': {'name': _c('插销（L 型盖）', 'Latch pin (L-Type cover)'),
                  'evidence': [_ev('material', CHECK_MATERIAL)]},
}


def extract_dn(code, description=''):
    """DN size from an RSV item code (RSV0080…, RSV010016…, RSVPE125…, RSVSO100…)
    or from 'DN100' in the description."""
    code = (code or '').upper().strip()
    m = re.match(r'^RSV0(\d{3})', code)
    if m:
        return int(m.group(1))
    m = re.match(r'^RSV(?:PE|SO|SP|CAP)(\d{2,3})', code)
    if m:
        return int(m.group(1))
    m = re.search(r'\bDN\s?(\d{2,3})\b', (description or '').upper())
    return int(m.group(1)) if m else None


def classify(code, description='', ref=None):
    """Return the product type key (see PRODUCT_TYPES) or None.
    ref: dict with 'category', 'sub_category', 'pc_category' from the
    reference table, or None."""
    code_u = (code or '').upper().strip()
    desc_u = (description or '').upper()
    ref = ref or {}
    cat = (ref.get('category') or '').lower()
    sub = (ref.get('sub_category') or '').lower()
    pcat = (ref.get('pc_category') or '').lower()
    if ref.get('description') and not desc_u:
        desc_u = ref['description'].upper()
    # "GAL" variant (galvanised bolts), but not "... Gal. Pipe" (clamp suits a gal pipe)
    is_gal = bool(re.search(r'\bGAL\b(?!\.?\s*PIPE)', desc_u))

    if 'comp covers' in cat:
        return None                                   # not configured yet
    if code_u.startswith('RSVCAP') or 'caps' in sub or 'valve caps' in pcat:
        return 'valve_cap'
    if code_u.startswith('RSV') or cat == 'gate valves' or 'rsv' in pcat:
        dn = extract_dn(code_u, desc_u)
        if dn is None:
            m = re.search(r'dn\s?(\d{2,3})', pcat)
            dn = int(m.group(1)) if m else None
        if dn is not None and dn >= 375:
            return 'rsv_375'
        if dn is not None and dn >= 200:
            return 'rsv_large'
        return 'rsv_small'
    if code_u.startswith('UMG') or 'gibault' in pcat or 'GIBAULT' in desc_u:
        return 'gibault'
    if code_u.startswith('UMC') or pcat.startswith('umc') or cat == 'couplings':
        return 'umc_gal' if (is_gal or re.search(r'G[LS]$', code_u)) else 'umc'
    if code_u.startswith('GPE') or cat == 'pe fittings' or 'pe restraint' in pcat or 'GRIPPA' in desc_u:
        return 'pe_grippa'
    if code_u.startswith('WPB') or 'pretap bush' in pcat:
        return 'pretap_bush'
    if cat in ('di fittings', 'pretaps') or pcat.startswith('di fittings') or code_u.startswith('WDFP'):
        if code_u.startswith(('DFBF',)) and cat == 'blank flanges':
            return 'blank_flange'
        return 'di_fitting'
    if cat == 'blank flanges' or pcat.startswith('blank flanges') or 'BLANK FLANGE' in desc_u:
        return 'blank_flange'
    if cat == 'repair clamps' or 'repair clamp' in pcat or 'CLAMP' in desc_u:
        return 'repair_clamp_gal' if is_gal else 'repair_clamp'
    if cat == 'spindles' or 'extension spindle' in pcat or 'EXTENSION SPINDLE' in desc_u:
        return 'spindle'
    if cat == 'ss straps' or 'ss straps' in pcat or code_u.startswith('APSSSR'):
        return 'ss_strap'
    if cat == 'gaskets' or 'gasket' in pcat or code_u.startswith('GASK'):
        return 'gasket'
    if cat == 'handwheels' or 'handwheel' in pcat or code_u.startswith('AVHW'):
        return 'handwheel'
    if code_u.startswith('ACLTYPELP') or 'LATCH PIN' in desc_u:
        return 'latch_pin'
    # hydrant heads (single / dual, CFA / MFB) are not covers
    if code_u.startswith(('ACLTYPES', 'ACLTYPED')) or (
            ('L - TYPE' in desc_u or 'L-TYPE' in desc_u) and 'HYDRANT HEAD' in desc_u):
        return 'l_type_head'
    if code_u.startswith('ACLTYPE') or 'l - type' in pcat or 'L - TYPE' in desc_u or 'L-TYPE' in desc_u:
        return 'l_type'
    if cat == 'covers & lids' or 'access covers' in pcat or 'fire plug' in pcat:
        return 'cover'
    # Fallbacks when the reference table does not know the code
    if code_u.startswith('DF'):
        return 'di_fitting'
    if code_u.startswith('AC'):
        return 'cover'
    return None


# ── Reference table import ───────────────────────────────────────────────────

def _norm(value):
    return ' '.join(str(value or '').split())


def parse_reference_workbook(data):
    """Parse reference.xlsx into {CODE: {...}}. Uses the 'ReferenceData' sheet
    (Item Code / Description / Category / Sub Category) and the 'Product
    Codes' sheet (Code / Description / Category) for the finer category."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = {}

    def rows_of(ws):
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None) or ()
        names = [_norm(h).lower() for h in header]
        for row in rows:
            yield {names[i]: row[i] for i in range(min(len(names), len(row)))}

    for ws in wb.worksheets:
        first = next(ws.iter_rows(values_only=True, max_row=1), ()) or ()
        heads = [_norm(h).lower() for h in first]
        if 'item code' in heads and 'sub category' in heads:          # ReferenceData
            for r in rows_of(ws):
                code = _norm(r.get('item code')).upper()
                if not code:
                    continue
                item = out.setdefault(code, {'code': code})
                item['description'] = _norm(r.get('description')) or item.get('description', '')
                item['category'] = _norm(r.get('category'))
                item['sub_category'] = _norm(r.get('sub category'))
        elif 'code' in heads and any(h.startswith('category') for h in heads):  # Product Codes
            cat_key = next(h for h in heads if h.startswith('category'))
            for r in rows_of(ws):
                code = _norm(r.get('code')).upper()
                if not code:
                    continue
                item = out.setdefault(code, {'code': code})
                item.setdefault('description', _norm(r.get('description')))
                item['pc_category'] = _norm(r.get(cat_key))
    wb.close()
    return out
