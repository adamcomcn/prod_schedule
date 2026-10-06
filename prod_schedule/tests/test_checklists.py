"""Checklist templates: Excel import, editor, versions and answer evaluation."""
import io
import json
import os
import sys
import tempfile
import unittest

os.environ.setdefault('APP_DATA_DIR', tempfile.mkdtemp(prefix='prod-schedule-tests-'))
os.environ.setdefault('SECRET_KEY', 'test-secret-key')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app
import checklists
from db import db_conn
from openpyxl import Workbook


def sample_workbook(with_chinese=True):
    """The Daemco checklist layout, cut down."""
    wb = Workbook()
    ws = wb.active
    ws['C2'] = 'DPL:'
    ws['C3'] = 'DUCTILE IRON FITTINGS - DAEMCO EXTERNAL INSPECTION CHECKLIST - V1.0'
    header = ['', '', 'PART', 'Q.', 'INSPECTION GUIDELINE', 'IF', 'WHAT TO DO?']
    if with_chinese:
        header += ['', '', '', '', 'PART 中文', 'INSPECTION GUIDELINE 中文', 'WHAT TO DO? 中文']
    ws.append([])
    ws.append(header)
    rows = [
        ('ASSEMBLY LEVEL CHECK', 'Product Marking', 1, 'Is the DN correct to the bore ID?', 'If no', 'Reject',
         '产品标识', '标注的 DN 是否与内孔尺寸一致？', '拒收'),
        ('', '', 2, 'What is the legibility of these markings? [Good/ Fair/ Poor]', 'If Poor', 'Reject',
         '', '标识清晰度', '拒收'),
        ('', 'External Body', 1, 'Are the feet scuffed? (IF APPLICABLE)', 'If yes', 'Rework',
         '外部本体', '底脚是否有擦伤？', '返工'),
        ('', '', 2, 'What is the coating thickness? (External > 300μm)', 'If < 300μm', 'Reject',
         '', '外涂层厚度', '拒收'),
        ('', 'Spigot (IF APPLICABLE)', 1, 'Are there any casting defects?', 'If yes', 'Reject',
         '插口', '是否有铸造缺陷？', '拒收'),
    ]
    for sect, part, num, text, cond, action, part_zh, text_zh, action_zh in rows:
        line = ['', sect, part, num, text, cond, action]
        if with_chinese:
            line += ['', '', '', '', part_zh, text_zh, action_zh]
        ws.append(line)
    ws.append(['', 'DATE'])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ParseTests(unittest.TestCase):
    def test_parse_layout_rules_and_chinese(self):
        t = checklists.parse_checklist_workbook(sample_workbook())
        self.assertIn('CHECKLIST', t['title'])
        self.assertEqual([s['name'] for s in t['sections']], ['Product Marking', 'External Body', 'Spigot'])
        self.assertTrue(t['sections'][2]['optional'])
        dn, legibility = t['sections'][0]['questions']
        self.assertEqual((dn['type'], dn['fail_on'], dn['text_zh'], dn['action_zh']),
                         ('yes_no', 'no', '标注的 DN 是否与内孔尺寸一致？', '拒收'))
        self.assertEqual((legibility['type'], legibility['fail_on']), ('rating', 'poor'))
        self.assertNotIn('[Good', legibility['text'])
        feet, thickness = t['sections'][1]['questions']
        self.assertTrue(feet['optional'])
        self.assertNotIn('IF APPLICABLE', feet['text'])
        self.assertEqual((thickness['type'], thickness['min'], thickness['unit']), ('number', 300, 'μm'))
        self.assertEqual(t['sections'][0]['name_zh'], '产品标识')

    def test_evaluate(self):
        yes_fail = {'type': 'yes_no', 'fail_on': 'yes'}
        self.assertEqual(checklists.evaluate(yes_fail, 'yes'), 'fail')
        self.assertEqual(checklists.evaluate(yes_fail, 'no'), 'ok')
        self.assertEqual(checklists.evaluate(yes_fail, 'na'), 'na')
        self.assertEqual(checklists.evaluate(yes_fail, ''), '')
        rating = {'type': 'rating', 'fail_on': 'poor'}
        self.assertEqual(checklists.evaluate(rating, 'fair'), 'ok')
        self.assertEqual(checklists.evaluate(dict(rating, fail_on='fair'), 'fair'), 'fail')
        number = {'type': 'number', 'min': 300}
        self.assertEqual(checklists.evaluate(number, '299.5'), 'fail')
        self.assertEqual(checklists.evaluate(number, '342'), 'ok')
        self.assertEqual(checklists.evaluate(number, 'abc'), '')

    def test_bad_files(self):
        wb = Workbook()
        wb.active.append(['nothing here'])
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(ValueError):
            checklists.parse_checklist_workbook(buf.getvalue())
        with self.assertRaises(ValueError):
            checklists.normalise_template(
                {'sections': [{'name': 'A', 'questions': [{'text': 'x', 'type': 'number'}]}]})


class PageTests(unittest.TestCase):
    def setUp(self):
        app.app.config.update(TESTING=True)
        with db_conn() as conn:
            for table in ('users', 'checklist_templates', 'checklist_versions'):
                conn.execute(f'DELETE FROM {table}')
            conn.execute("INSERT INTO users (username,password_hash,role) VALUES ('murphy','x','admin')")
            conn.execute("INSERT INTO users (username,password_hash,role) VALUES ('yu','x','inspector')")
            self.ids = {r['username']: r['id'] for r in conn.execute('SELECT id, username FROM users')}

    def client_for(self, name):
        client = app.app.test_client()
        with client.session_transaction() as session:
            session['user_id'] = self.ids[name]
            session['_csrf_token'] = 'tok'
        return client

    def import_sample(self, client):
        return client.post('/checklists/import', data={
            '_csrf_token': 'tok', 'target': 'new', 'name': 'DI Fittings',
            'file': (io.BytesIO(sample_workbook()), 'di.xlsx')}, content_type='multipart/form-data')

    def test_import_edit_versions_and_lookup(self):
        murphy = self.client_for('murphy')
        self.assertEqual(self.import_sample(murphy).status_code, 302)
        with db_conn() as conn:
            tid = conn.execute('SELECT id FROM checklist_templates').fetchone()[0]
        _tpl, ver, data = app.checklist_version(tid)
        self.assertEqual(ver['version'], 1)
        self.assertIn('Is the DN correct to the bore ID?', murphy.get(f'/checklists/{tid}/edit').get_data(as_text=True))
        self.assertIsNone(app.checklist_for_product_type('di_fitting'))      # no product type yet

        # fix a condition, add Chinese, assign the product type -> version 2
        data['sections'][1]['questions'][0]['fail_on'] = 'no'
        data['sections'][0]['questions'][0]['text_zh'] = '内孔 DN 是否一致？'
        murphy.post(f'/checklists/{tid}/edit', data={
            '_csrf_token': 'tok', 'name': 'DI Fittings', 'active': '1', 'product_types': ['di_fitting', 'bogus'],
            'note': 'feet condition', 'data_json': json.dumps(data)})
        tpl, ver, data2 = app.checklist_version(tid)
        self.assertEqual((ver['version'], ver['note'], tpl['product_types']), (2, 'feet condition', 'di_fitting'))
        self.assertEqual(data2['sections'][1]['questions'][0]['fail_on'], 'no')
        self.assertEqual(data2['sections'][0]['questions'][0]['id'], data['sections'][0]['questions'][0]['id'])
        self.assertEqual(app.checklist_version(tid, 1)[2]['sections'][1]['questions'][0]['fail_on'], 'yes')
        self.assertEqual(app.checklist_for_product_type('di_fitting')['id'], tid)

        # saving unchanged questions does not create a version
        murphy.post(f'/checklists/{tid}/edit', data={'_csrf_token': 'tok', 'name': 'DI Fittings', 'active': '1',
                                                    'product_types': 'di_fitting', 'data_json': json.dumps(data2)})
        self.assertEqual(app.checklist_version(tid)[1]['version'], 2)

        # invalid data is refused
        murphy.post(f'/checklists/{tid}/edit', data={'_csrf_token': 'tok', 'name': 'x', 'active': '1',
                                                    'data_json': json.dumps({'sections': []})})
        self.assertEqual(app.checklist_version(tid)[1]['version'], 2)

        self.assertIn('标注的 DN', murphy.get(f'/checklists/{tid}/preview?v=1').get_data(as_text=True))
        reimported = checklists.parse_checklist_workbook(murphy.get(f'/checklists/{tid}/export.xlsx').get_data())
        self.assertEqual(checklists.question_count(reimported), 5)
        self.assertEqual(reimported['sections'][1]['questions'][0]['fail_on'], 'no')
        self.assertEqual(reimported['sections'][0]['questions'][0]['text_zh'], '内孔 DN 是否一致？')

    def test_inspectors_cannot_edit(self):
        yu = self.client_for('yu')
        self.assertEqual(yu.get('/checklists').status_code, 403)
        self.assertEqual(self.import_sample(yu).status_code, 403)
        self.assertNotIn('/checklists', yu.get('/tasks').get_data(as_text=True))

    def test_only_admins_see_checklists(self):
        with db_conn() as conn:
            conn.execute("INSERT INTO users (username,password_hash,role) VALUES ('lead2','x','lead')")
            self.ids['lead2'] = conn.execute("SELECT id FROM users WHERE username='lead2'").fetchone()[0]
        lead = self.client_for('lead2')
        self.assertEqual(lead.get('/checklists').status_code, 403)
        self.assertEqual(lead.get('/checklists/bulk').status_code, 403)
        self.assertNotIn('/checklists', lead.get('/tasks').get_data(as_text=True))
        self.assertIn('/checklists', self.client_for('murphy').get('/tasks').get_data(as_text=True))   # admin


if __name__ == '__main__':
    unittest.main()
