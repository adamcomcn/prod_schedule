"""Required-evidence rules (BRT_required_evidence.pdf + agreed decisions)."""
import io
import os
import sys
import tempfile
import unittest

os.environ.setdefault('APP_DATA_DIR', tempfile.mkdtemp(prefix='prod-schedule-tests-'))
os.environ.setdefault('SECRET_KEY', 'test-secret-key')

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app
import evidence_rules as er
from db import db_conn
from helpers import assign_job
from openpyxl import Workbook


def ref(category='', sub='', pc=''):
    return {'category': category, 'sub_category': sub, 'pc_category': pc}


def types(ptype):
    return [e['type'] for e in er.PRODUCT_TYPES[ptype]['evidence']]


class ClassifyTests(unittest.TestCase):
    def test_agreed_rules(self):
        cases = [
            # (code, description, reference, expected type)
            ('RSV0100FLFLCC', 'DN100 Resilient Seated Gate Valve', ref('Gate Valves', 'Gate Valves', 'FL RSV - DN100'), 'rsv_small'),
            ('RSVPE180ACC', 'DN180 Resilient Seated Gate Valve ACC PE PE', ref('Gate Valves'), 'rsv_small'),   # PE follows RSV
            ('RSVSO150ACC', 'DN150 Socket Gate Valve', None, 'rsv_small'),
            ('RSV022516FLFL', 'DN225 Resilient Seated Gate Valve', None, 'rsv_large'),
            ('RSVSO200ACC', 'DN200 Socket Gate Valve', None, 'rsv_large'),                                    # spark for all types
            ('RSV0375FLFLCC', 'DN375 Resilient Seated Gate Valve', None, 'rsv_375'),
            ('RSVCAP100', 'Valve cap', ref('Gate Valves', 'Gate Valves - Caps'), 'valve_cap'),
            ('UMC0100316L', 'DN100 Long Unrestrained Mechanical Couplings', ref('Couplings'), 'umc'),
            ('UMC0100GL', 'DN100 Long Unrestrained Mechanical Couplings GAL', ref('Couplings'), 'umc_gal'),
            ('UMG080PVCPVC', 'GIBAULT JOINT COUPLING DN80', ref(pc='Gibault'), 'gibault'),
            ('DFT1008F', 'DN100 x 80 Flange Tee DI', ref('DI Fittings', 'DI Fittings - Flange Tee'), 'di_fitting'),
            ('WDFPQ1525', 'DN150 Pretap Connector Quad', ref('Pretaps', 'DI Fittings - Pretap'), 'di_fitting'),
            ('wPB25', 'Pretap Bush Suit 25mm', ref('DI Fittings', 'DI Fittings - Pretap', 'Pretap Bush & O-Rings'), 'pretap_bush'),
            ('DFBF010016', 'DI BLANK FLANGE DN100', ref('Blank Flanges'), 'blank_flange'),
            ('GPESOC125', 'Grippa PE Socket Coupling', ref('PE Fittings'), 'pe_grippa'),
            ('NCR7583-150', 'NC REPAIR CLAMP', ref('Repair Clamps'), 'repair_clamp'),
            ('NCFCG20-75', 'Junior Clamps Full Circle Gasket Gal. Pipe', ref('Repair Clamps'), 'repair_clamp'),
            ('ES1000', 'Extension Spindle 1000', ref('Spindles'), 'spindle'),
            ('APSSSR0200', 'SS316 STRAP', ref('SS Straps'), 'ss_strap'),
            ('GASKB0080TD', 'DN80 BLUE EPDM GASKET', None, 'gasket'),                                         # prefix fallback
            ('AVHW100180CC', 'Handwheel', ref(pc='Accessories - Handwheels'), 'handwheel'),
            ('ACTBW', 'DI Toby Box W', ref('Covers & Lids'), 'cover'),
            ('ACLTYPE', 'L-type hydrant cover', ref('Covers & Lids'), 'l_type'),
            ('ACLTYPESMFB', 'Single MFB L - Type Hydrant Head', ref(pc='L - Type Hydrant'), 'l_type_head'),
            ('ACLTYPEDCFA', 'Dual CFA L - Type Hydrant Head', ref('Covers & Lids'), 'l_type_head'),
            ('XYZ1', 'Dual MFB L - Type Hydrant Head', None, 'l_type_head'),
            ('ACLTYPESC', 'L - Type Hydrant Cover Surround Concrete', None, 'l_type'),
            ('WAPVAL', 'Valve Anchor Legs (Suit DN100 - DN150)', ref('DI Fittings'), 'valve_legs'),
            ('ACLTYPELP', 'Latch Pin (Suit L Type Cover)', ref('Covers & Lids'), 'latch_pin'),
            ('ACCOMP1', 'Composite cover', ref('Comp Covers', 'Comp Covers - SMC Lid'), None),               # left out
        ]
        for code, desc, reference, expected in cases:
            with self.subTest(code=code):
                self.assertEqual(er.classify(code, desc, reference), expected)

    def test_required_evidence_per_type(self):
        self.assertEqual(types('rsv_small'), ['brt', 'daq', 'vtrust'])
        self.assertEqual(types('rsv_large'), ['brt', 'spark', 'daq', 'vtrust'])
        self.assertEqual(types('rsv_375'), ['brt', 'spark', 'vtrust'])          # no DAQ
        self.assertEqual(types('di_fitting'), ['brt', 'pressure'])
        self.assertEqual(types('pretap_bush'), ['brt'])                         # no pressure test
        self.assertEqual(types('repair_clamp'), ['brt', 'xrf'])
        self.assertEqual(types('l_type'), ['checklist'])
        self.assertEqual(types('l_type_head'), ['checklist'])
        self.assertEqual(types('latch_pin'), ['material'])
        self.assertEqual(types('cover'), ['checklist', 'material'])
        self.assertIn('Q235B', er.PRODUCT_TYPES['umc_gal']['evidence'][0]['checks'][-1]['en'])

    def test_parse_reference_workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Product Codes'
        ws.append(['Code', 'Description', 'Category (For NC CA PA Table only, not from Xero)'])
        ws.append(['UMG080PVCPVC', 'GIBAULT JOINT COUPLING DN80', 'Gibault'])
        ws.append(['RSV0100FLFLCC', 'DN100 RSV', 'FL RSV - DN100'])
        ws2 = wb.create_sheet('ReferenceData')
        ws2.append(['Item Code', 'Description', 'Category', 'Sub Category'])
        ws2.append(['RSV0100FLFLCC', 'DN100 Resilient Seated Gate Valve', 'Gate Valves', 'Gate Valves'])
        buf = io.BytesIO()
        wb.save(buf)
        items = er.parse_reference_workbook(buf.getvalue())
        self.assertEqual(items['RSV0100FLFLCC']['category'], 'Gate Valves')
        self.assertEqual(items['RSV0100FLFLCC']['pc_category'], 'FL RSV - DN100')
        self.assertEqual(items['UMG080PVCPVC']['pc_category'], 'Gibault')


class EvidenceFlowTests(unittest.TestCase):
    JOB = 'MELBOURNE|PO-5|RSV0100FLFLCC'

    def setUp(self):
        app.app.config.update(TESTING=True)
        with db_conn() as conn:
            for table in ('users', 'inspection_tasks', 'inspection_attachments', 'product_reference'):
                conn.execute(f'DELETE FROM {table}')
            conn.execute("INSERT INTO users (username,password_hash,role,display_name) VALUES ('boss','x','admin','Boss')")
            conn.execute("INSERT INTO users (username,password_hash,role,display_name) VALUES ('yu','x','inspector','Yu')")
            self.ids = {r['username']: r['id'] for r in conn.execute('SELECT id, username FROM users')}
        schedule = {'MELBOURNE': [['Order Number', 'Daemco Purchase Order', 'Item Code', 'Item Description', 'Quantity'],
                                  ['DPL5', 'PO-5', 'RSV0100FLFLCC', 'DN100 Resilient Seated Gate Valve', '10'],
                                  ['DPL6', 'PO-6', 'XYZ999', 'Mystery product', '1']]}
        app.save_json(app.CURRENT_FILE, schedule)
        app.save_json(app.PREVIOUS_FILE, schedule)
        app.save_json(app.INSPECTIONS_CACHE, {})
        config = app.load_config()
        config.update(hq_report_emails='', task_notify_emails='', modules={})
        app.save_json(app.CONFIG_FILE, config)
        assign_job(self.JOB, 'yu')

    def client_for(self, name):
        client = app.app.test_client()
        with client.session_transaction() as session:
            session['user_id'] = self.ids[name]
            session['_csrf_token'] = 'tok'
        return client

    def test_inspection_page_shows_checks_and_daq(self):
        page = self.client_for('yu').get(f'/inspect/{self.JOB}').get_data(as_text=True)
        self.assertIn('name="ev_check_brt_2"', page)          # 3 BRT check items
        self.assertIn('name="daq_t3"', page)
        self.assertIn('data-limit="2.4"', page)
        self.assertNotIn('data-ev="spark"', page)             # DN100: no spark test
        self.assertIn('id="ev-cam-brt" accept="image/*" capture="environment"', page)  # phone camera
        self.assertIn('id="ev-vid-vtrust" accept="video/*"', page)

    def test_submit_records_checks_daq_and_missing(self):
        yu = self.client_for('yu')
        yu.post(f'/inspect/{self.JOB}/submit', data={
            '_csrf_token': 'tok', 'region': 'MELBOURNE', 'order_number': 'DPL5', 'item_code': 'RSV0100FLFLCC',
            'item_description': 'DN100 Resilient Seated Gate Valve', 'inspection_date': '2026-10-09',
            'quantity_inspected': '10', 'result': 'Fail',
            'ev_result_brt': 'Fail', 'ev_check_brt_0': 'ok', 'ev_check_brt_1': 'fail',
            'ev_file_brt': (io.BytesIO(b'%PDF brt'), 'brt.pdf'),
            'ev_result_daq': 'Fail', 'daq_t1': '1.80', 'daq_t2': '1.70', 'daq_t3': '2.5',
            'ev_result_vtrust': '',
        }, content_type='multipart/form-data')
        record = app.load_json(app.INSPECTIONS_CACHE)[self.JOB][0]
        self.assertEqual(record['product_type'], 'rsv_small')
        brt = record['evidence']['brt']
        self.assertEqual([c['state'] for c in brt['checks']], ['ok', 'fail', ''])
        self.assertEqual(record['evidence']['daq']['daq_values'], {'t1': 1.8, 't2': 1.7, 't3': 2.5})
        self.assertFalse(record['evidence']['daq']['daq_ok'])  # T2 below 1.76
        self.assertEqual(record['missing_evidence'], ['daq', 'vtrust'])  # no files for them
        pdf = self.client_for('yu').get(f'/inspect/{self.JOB}/report.pdf')
        self.assertEqual(pdf.status_code, 200)

    def test_reference_upload_and_coverage(self):
        boss = self.client_for('boss')
        page = boss.get('/settings').get_data(as_text=True)
        self.assertIn('XYZ999', page)                         # unknown product listed
        wb = Workbook()
        ws = wb.active
        ws.title = 'ReferenceData'
        ws.append(['Item Code', 'Description', 'Category', 'Sub Category'])
        ws.append(['XYZ999', 'Mystery product', 'Spindles', 'Spindles'])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = boss.post('/settings/reference', data={'_csrf_token': 'tok', 'reference': (buf, 'reference.xlsx')},
                             content_type='multipart/form-data')
        self.assertEqual(response.status_code, 302)
        self.assertEqual(app.product_type_for('XYZ999', 'Mystery product'), 'spindle')
        self.assertNotIn('<code>XYZ999</code>', boss.get('/settings').get_data(as_text=True))
        # inspectors cannot import
        self.assertEqual(self.client_for('yu').post('/settings/reference', data={'_csrf_token': 'tok'}).status_code, 403)


if __name__ == '__main__':
    unittest.main()
