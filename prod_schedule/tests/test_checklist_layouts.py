"""The other Daemco checklist layouts: Chinese and English in one cell,
'Item / If no, what to do?', sub-part columns, sentence conditions."""
import io
import os
import sys
import tempfile
import unittest

os.environ.setdefault('APP_DATA_DIR', tempfile.mkdtemp(prefix='prod-schedule-tests-'))
os.environ.setdefault('SECRET_KEY', 'test-secret-key')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import checklists
from openpyxl import Workbook


def workbook(header, rows, title=None, header_row=4):
    wb = Workbook()
    ws = wb.active
    if title:
        ws.cell(row=2, column=3, value=title)
    for col, value in enumerate(header, 1):
        if value:
            ws.cell(row=header_row, column=col, value=value)
    for r, row in enumerate(rows, header_row + 1):
        for col, value in enumerate(row, 1):
            if value not in (None, ''):
                ws.cell(row=r, column=col, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class SplitBilingualTests(unittest.TestCase):
    def test_cases(self):
        cases = {
            '部件是否用黑色沥青涂装？Is the part coated in black bitumen?': ('Is the part coated in black bitumen?', '部件是否用黑色沥青涂装？'),
            'Can an M16 bolt fit into both the bolt holes?M16螺栓是否适配两个螺栓孔？':
                ('Can an M16 bolt fit into both the bolt holes?', 'M16螺栓是否适配两个螺栓孔？'),
            'Is the marking correct (DAEMCO)?标识DAEMCO是否正确？': ('Is the marking correct (DAEMCO)?', '标识DAEMCO是否正确？'),
            '组装时盖子是否平齐Is the lid flush?\n盖子上方最大距离2mm': ('Is the lid flush?', '组装时盖子是否平齐 盖子上方最大距离2mm'),
            '水盒盖Lid': ('Lid', '水盒盖'),
            '返工或拒收Rework or Reject': ('Rework or Reject', '返工或拒收'),
            'Plain English only': ('Plain English only', ''),
            '只有中文': ('', '只有中文'),
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(checklists.split_bilingual(text), expected)


class LayoutTests(unittest.TestCase):
    def test_bilingual_cells_and_headers(self):
        data = workbook(['', '', 'PART', 'Q.', '检查指引INSPECTION GUIDELINE', '如果IF', '如何做WHAT TO DO?'], [
            ['', '装配水平检查 ASSEMBLY', '水盒盖Lid', 1, '部件是否用黑色沥青涂装？Is the part coated in black bitumen?',
             '如果不是If no', '返工或拒收Rework or Reject'],
            ['', '', '', 2, '评估零件的清洁度：[良好、一般、较差]Rate the cleanliness of the part: [Good, Fair, Poor]',
             '如果较差If poor', '清洁或拒收Clean or Reject'],
            ['', '', '水盒体Frame', 1, '有锋利的边缘吗？Is there any sharp edges?', '如果是If yes', '打磨锋利边角Break all sharp edges'],
            ['', '日期DATE'],
        ], title='199水盒盖- DAEMCO 检查确认单 - V1.0')
        t = checklists.parse_checklist_workbook(data)
        self.assertEqual([(s['name'], s['name_zh']) for s in t['sections']], [('Lid', '水盒盖'), ('Frame', '水盒体')])
        coat, clean = t['sections'][0]['questions']
        self.assertEqual((coat['text'], coat['text_zh'], coat['fail_on'], coat['action'], coat['action_zh']),
                         ('Is the part coated in black bitumen?', '部件是否用黑色沥青涂装？', 'no', 'Rework or Reject', '返工或拒收'))
        self.assertEqual((clean['type'], clean['fail_on']), ('rating', 'poor'))
        self.assertNotIn('[', clean['text'] + clean['text_zh'])
        self.assertEqual(t['sections'][1]['questions'][0]['fail_on'], 'yes')

    def test_item_column_and_if_no_header(self):
        data = workbook(['', 'Item', 'Q.', 'Inspection Guideline', 'Yes or No', 'If no, what to do?'], [
            ['', 'Lid', 1, 'Is the part free from any sharp edges?', '', 'Break all sharp edges'],
            ['', '', 2, 'Is the lid marked "DAEMCO"?', '', 'Reject'],
            ['', 'Frame', 1, 'Is the Frame marked "MMYY"?', '', 'Reject'],
            ['', 'Production Batch:'], ['', 'Inspection Date:'], ['', 'Signature:'],
        ], header_row=3)
        t = checklists.parse_checklist_workbook(data)
        self.assertEqual([s['name'] for s in t['sections']], ['Lid', 'Frame'])
        self.assertTrue(all(q['fail_on'] == 'no' for s in t['sections'] for q in s['questions']))
        self.assertEqual(t['sections'][0]['questions'][0]['action'], 'Break all sharp edges')

    def test_sub_part_column_less_than_sentence_and_errors(self):
        data = workbook(['', 'PART', '', 'Q.', 'INSPECTION GUIDELINE', 'IF', 'WHAT TO DO?'], [
            ['', 'COMPONENTS', 'Coated Valve Cap', 1, 'What is the coating thickness? (External > 300 μm)', 'If less than', 'Reject'],
            ['', '', '', 2, 'Can the eyelets screw down properly? [If Applicable]', 'If no', 'Re-tap hole or Reject'],
            ['', '', '', '', '#VALUE!'],
            ['', 'ASSEMBLY', 'Evidence', 1, 'Before pressure testing – Provide a video showing the gasket installed',
             'Reject the valve if the gasket does not fit properly.', ''],
            ['', '', '', 2, 'What is the coating thickness ? (External > 300μm)', 'If no', 'Rework'],
            ['', 'DATE'],
        ])
        t = checklists.parse_checklist_workbook(data)
        self.assertEqual([s['name'] for s in t['sections']], ['Coated Valve Cap', 'Evidence'])
        thickness, eyelets = t['sections'][0]['questions']
        self.assertEqual((thickness['type'], thickness['min'], thickness['unit']), ('number', 300, 'μm'))
        self.assertTrue(eyelets['optional'])
        video, measure = t['sections'][1]['questions']
        self.assertEqual((video['fail_on'], video['photo'], video['action']),
                         ('no', 'always', 'Reject the valve if the gasket does not fit properly.'))
        self.assertEqual((measure['type'], measure['min']), ('number', 300))


if __name__ == '__main__':
    unittest.main()
