import io
import os
import sys
import tempfile
import unittest
import openpyxl

TEST_DATA_DIR = tempfile.mkdtemp(prefix='prod-schedule-tests-')
os.environ['APP_DATA_DIR'] = TEST_DATA_DIR
os.environ['SECRET_KEY'] = 'test-secret-key'

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app
from db import db_conn
from openpyxl import Workbook
from werkzeug.security import generate_password_hash


class SecurityAndProductionTests(unittest.TestCase):
    def setUp(self):
        app.app.config.update(TESTING=True)
        app._login_failures.clear()
        with db_conn() as conn:
            conn.execute('DELETE FROM users')
            conn.execute('DELETE FROM outstanding_jobs')
            conn.execute('DELETE FROM schedule_uploads')
            conn.execute(
                'INSERT INTO users (username,password_hash,role) VALUES (?,?,?)',
                ('admin-test', generate_password_hash('admin-test-password'), 'admin'))
            conn.execute(
                'INSERT INTO users (username,password_hash,role) VALUES (?,?,?)',
                ('inspector-test', generate_password_hash('inspector-test-password'), 'inspector'))
        self.client = app.app.test_client()

    def login_as(self, username):
        with db_conn() as conn:
            user_id = conn.execute(
                'SELECT id FROM users WHERE username=?', (username,)
            ).fetchone()[0]
        with self.client.session_transaction() as session:
            session['user_id'] = user_id
            session['_csrf_token'] = 'test-csrf-token'

    def test_health_is_public_and_business_routes_require_login(self):
        self.assertEqual(self.client.get('/healthz').status_code, 200)
        self.assertEqual(self.client.get('/').status_code, 302)
        self.assertIn('/login', self.client.get('/').location)

    def test_post_requires_csrf(self):
        response = self.client.post(
            '/login',
            data={'username': 'admin-test', 'password': 'admin-test-password'})
        self.assertEqual(response.status_code, 400)

    def test_inspector_cannot_access_admin_endpoints(self):
        self.login_as('inspector-test')
        self.assertEqual(self.client.get('/settings').status_code, 403)
        self.assertEqual(self.client.get('/admin/users').status_code, 403)
        response = self.client.post(
            '/upload',
            data={'_csrf_token': 'test-csrf-token'},
            content_type='multipart/form-data')
        self.assertEqual(response.status_code, 403)

    def test_admin_can_manage_users_without_default_passwords(self):
        self.login_as('admin-test')
        self.assertEqual(self.client.get('/admin/users').status_code, 200)
        with db_conn() as conn:
            usernames = {row[0] for row in conn.execute('SELECT username FROM users')}
        self.assertNotIn('admin', usernames)
        self.assertNotIn('qc1', usernames)
        self.assertNotIn('qc2', usernames)

    def test_upload_rejects_non_xlsx_and_security_headers_are_set(self):
        self.login_as('admin-test')
        response = self.client.post(
            '/upload',
            data={
                '_csrf_token': 'test-csrf-token',
                'file': (io.BytesIO(b'not an excel file'), 'payload.txt'),
            },
            content_type='multipart/form-data')
        self.assertEqual(response.status_code, 302)
        page = self.client.get('/')
        self.assertEqual(page.headers['X-Frame-Options'], 'DENY')
        self.assertEqual(page.headers['X-Content-Type-Options'], 'nosniff')

    def test_excel_parser_accepts_plain_xlsx_and_rejects_invalid_content(self):
        workbook = Workbook()
        workbook.active.append(['Order Number', 'Item Code'])
        workbook.active.append(['PO-1', 'ITEM-1'])
        content = io.BytesIO()
        workbook.save(content)

        parsed = app.parse_excel(content.getvalue(), '')
        self.assertEqual(parsed['Sheet'][1], ['PO-1', 'ITEM-1'])

        with self.assertRaises(app.InvalidExcelFile):
            app.parse_excel(b'not an excel file', '')

    def test_encrypted_excel_without_password_has_clear_error(self):
        encrypted_office_header = bytes.fromhex('D0CF11E0A1B11AE1') + (b'\0' * 512)
        with self.assertRaises(app.ExcelPasswordRequired):
            app.decrypt_excel(encrypted_office_header, '')

    def test_all_fully_shipped_jobs_are_retained_and_require_qa_brt(self):
        headers = [
            'Daemco Purchase Order', 'Item Code', 'Item Description',
            'Supplier', 'Quantity', 'QA BRTs Sent?']
        previous = {
            'MELBOURNE': [
                headers,
                ['PO-DONE', 'ITEM-1', 'Done item', 'Supplier A', '2', 'YES'],
                ['PO-PENDING', 'ITEM-2', 'Pending item', 'Supplier B', '3', 'NO'],
            ]}
        current = {'MELBOURNE': [headers]}
        _, _, shipped_rows = app.compute_changes(previous, current)

        persisted, pending = app.persist_fully_shipped_jobs(
            previous, current, shipped_rows, '12 Jun 2026 12:00')

        self.assertEqual(persisted, 2)
        self.assertEqual(pending, 1)
        with db_conn() as conn:
            rows = conn.execute(
                'SELECT job_key, completed FROM outstanding_jobs ORDER BY job_key'
            ).fetchall()
        self.assertEqual(len(rows), 2)
        self.assertEqual(
            {row['job_key']: row['completed'] for row in rows},
            {
                'MELBOURNE|PO-DONE|ITEM-1': 1,
                'MELBOURNE|PO-PENDING|ITEM-2': 0,
            })

        app.save_json(app.CURRENT_FILE, current)
        app.save_json(app.PREVIOUS_FILE, current)
        self.login_as('admin-test')
        response = self.client.get('/inspect/MELBOURNE|PO-PENDING|ITEM-2')
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Pending item', response.data)

    def test_upload_blocks_saved_previous_week_rollback(self):
        headers = ['Order Number', 'Item Code', 'Item Description']
        previous = {'MELBOURNE': [headers, ['OLD-PO', 'OLD-ITEM', 'Old week']]}
        current = {'MELBOURNE': [headers, ['NEW-PO', 'NEW-ITEM', 'New week']]}
        app.save_json(app.PREVIOUS_FILE, previous)
        app.save_json(app.CURRENT_FILE, current)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'MELBOURNE'
        for row in previous['MELBOURNE']:
            worksheet.append(row)
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)

        self.login_as('admin-test')
        response = self.client.post(
            '/upload',
            data={
                '_csrf_token': 'test-csrf-token',
                'file': (content, 'old-week.xlsx'),
            },
            content_type='multipart/form-data',
            follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn('已阻止上传'.encode(), response.data)
        self.assertEqual(app.load_json(app.CURRENT_FILE, {}), current)

    def test_upload_blocks_any_previously_uploaded_schedule(self):
        headers = ['Order Number', 'Item Code', 'Item Description']
        historical = {'MELBOURNE': [headers, ['HIST-PO', 'HIST-ITEM', 'Historical']]}
        app.remember_schedule_upload(historical, '01 Jun 2026 09:00')
        app.save_json(
            app.PREVIOUS_FILE,
            {'MELBOURNE': [headers, ['PREV-PO', 'PREV-ITEM', 'Previous']]})
        current = {'MELBOURNE': [headers, ['CURR-PO', 'CURR-ITEM', 'Current']]}
        app.save_json(app.CURRENT_FILE, current)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'MELBOURNE'
        for row in historical['MELBOURNE']:
            worksheet.append(row)
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)

        self.login_as('admin-test')
        response = self.client.post(
            '/upload',
            data={
                '_csrf_token': 'test-csrf-token',
                'file': (content, 'historical-week.xlsx'),
            },
            content_type='multipart/form-data',
            follow_redirects=True)
        self.assertIn('之前已上传过'.encode(), response.data)
        self.assertEqual(app.load_json(app.CURRENT_FILE, {}), current)

    def test_status_legend_filters_schedule_rows(self):
        headers = ['Order Number', 'Item Code', 'Item Description', 'QA BRTs Sent?']
        previous = {
            'MELBOURNE': [
                headers,
                ['WAIT-PO', 'WAIT-ITEM', 'Waiting row', 'NO'],
                ['SHIP-PO', 'SHIP-ITEM', 'Shipped row', 'NO'],
            ]}
        current = {
            'MELBOURNE': [
                headers,
                ['WAIT-PO', 'WAIT-ITEM', 'Waiting row', 'NO'],
                ['NEW-PO', 'NEW-ITEM', 'New row', 'NO'],
            ]}
        app.save_json(app.PREVIOUS_FILE, previous)
        app.save_json(app.CURRENT_FILE, current)
        self.login_as('admin-test')

        new_response = self.client.get('/?sheet=MELBOURNE&status=new')
        self.assertIn(b'New row', new_response.data)
        self.assertNotIn(b'Waiting row', new_response.data)
        self.assertIn('清除筛选'.encode(), new_response.data)

        shipped_response = self.client.get('/?sheet=MELBOURNE&view=all&status=shipped')
        self.assertIn(b'Shipped row', shipped_response.data)
        self.assertNotIn(b'New row', shipped_response.data)

        no_qa_response = self.client.get(
            '/?sheet=MELBOURNE&view=all&status=no_qa_brt')
        self.assertIn(b'Shipped row', no_qa_response.data)
        self.assertNotIn(b'Waiting row', no_qa_response.data)

        self.assertEqual(
            self.client.get('/?sheet=MELBOURNE&view=all&status=vtrust').status_code,
            200)

    def test_comparison_export_contains_status_and_fully_shipped_history(self):
        headers = ['Order Number', 'Item Code', 'Item Description', 'QA BRTs Sent?']
        previous = {
            'MELBOURNE': [
                headers,
                ['SHIP-PO', 'SHIP-ITEM', 'Shipped row', 'NO'],
            ]}
        current = {'MELBOURNE': [headers]}
        app.save_json(app.PREVIOUS_FILE, previous)
        app.save_json(app.CURRENT_FILE, current)
        _, _, shipped_rows = app.compute_changes(previous, current)
        app.persist_fully_shipped_jobs(previous, current, shipped_rows, '12 Jun 2026')
        self.login_as('admin-test')

        response = self.client.get('/export/comparison.xlsx')
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response.content_type)

        workbook = openpyxl.load_workbook(io.BytesIO(response.data), read_only=True)
        self.assertIn('MELBOURNE', workbook.sheetnames)
        self.assertIn('Fully Shipped History', workbook.sheetnames)
        region_rows = list(workbook['MELBOURNE'].iter_rows(values_only=True))
        self.assertEqual(region_rows[0][:2], ('Comparison Status', 'QA BRT Alert'))
        self.assertIn(('SHIPPED', 'No QA BRT'), [row[:2] for row in region_rows[1:]])
        history_rows = list(
            workbook['Fully Shipped History'].iter_rows(values_only=True))
        self.assertIn(('SHIPPED', 'No QA BRT'), [row[:2] for row in history_rows[1:]])

    def test_schedule_is_paginated_and_knowledge_tolerates_bad_image_json(self):
        headers = ['Order Number', 'Item Code', 'Description']
        rows = [headers] + [[f'PO-{i}', f'ITEM-{i}', f'Description {i}'] for i in range(500)]
        app.save_json(app.CURRENT_FILE, {'MELBOURNE': rows})
        app.save_json(app.PREVIOUS_FILE, {'MELBOURNE': rows})
        with db_conn() as conn:
            conn.execute(
                "UPDATE kb_articles SET images='invalid-json' "
                "WHERE id=(SELECT MIN(id) FROM kb_articles)")
        self.login_as('admin-test')
        response = self.client.get('/?sheet=MELBOURNE&per_page=50')
        self.assertEqual(response.status_code, 200)
        self.assertLess(len(response.data), 200_000)
        config = app.load_config()
        config['modules'] = {'knowledge': True}
        app.save_json(app.CONFIG_FILE, config)
        try:
            self.assertEqual(self.client.get('/knowledge').status_code, 200)
        finally:
            config['modules'] = {}
            app.save_json(app.CONFIG_FILE, config)


if __name__ == '__main__':
    unittest.main()
