"""Inspection report list with filters and the Excel export."""
import io
import os
import sys
import tempfile
import unittest

os.environ.setdefault('APP_DATA_DIR', tempfile.mkdtemp(prefix='prod-schedule-tests-'))
os.environ.setdefault('SECRET_KEY', 'test-secret-key')

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app
from db import db_conn
from openpyxl import load_workbook


def record(day, result, inspector, supplier='', **extra):
    return {'inspection_date': day, 'submitted_at': day + 'T01:00:00', 'order_number': 'DPL1',
            'item_code': 'RSV0100', 'item_description': 'DN100 valve', 'result': result,
            'inspector_name': inspector, 'supplier': supplier, **extra}


class ReportsTests(unittest.TestCase):
    def setUp(self):
        app.app.config.update(TESTING=True)
        with db_conn() as conn:
            for table in ('users', 'inspection_reviews', 'inspection_tasks'):
                conn.execute(f'DELETE FROM {table}')
            conn.execute("INSERT INTO users (username,password_hash,role) VALUES ('hq','x','hq')")
            conn.execute("INSERT INTO users (username,password_hash,role) VALUES ('yu','x','inspector')")
            self.ids = {r['username']: r['id'] for r in conn.execute('SELECT id, username FROM users')}
            conn.execute("INSERT INTO inspection_tasks (job_key, region, order_number, item_code, status, supplier) "
                         "VALUES ('MELBOURNE|PO-2|B', 'MELBOURNE', 'DPL2', 'B', 'Completed', 'Valve Co')")
            conn.execute("INSERT INTO inspection_reviews (job_key, insp_index, status, reviewer_name, reviewed_at) "
                         "VALUES ('MELBOURNE|PO-1|A', 1, 'approved', 'Murphy', '2026-10-05 02:00')")
        app.save_json(app.INSPECTIONS_CACHE, {
            'MELBOURNE|PO-1|A': [record('2026-10-01', 'Fail', 'Mr. Yu', 'Casting Ltd', notes='=HYPERLINK("x")'),
                                 record('2026-10-04', 'Pass', 'Murphy', 'Casting Ltd')],
            'MELBOURNE|PO-2|B': [record('2026-10-03', 'Pass', 'Mr. Yu', missing_evidence=['daq'])],
        })

    def client_for(self, name):
        client = app.app.test_client()
        with client.session_transaction() as session:
            session['user_id'] = self.ids[name]
        return client

    def test_rows_supplier_fallback_and_review(self):
        rows = app.inspection_report_rows()
        self.assertEqual([r['date'] for r in rows], ['2026-10-04', '2026-10-03', '2026-10-01'])
        self.assertEqual(rows[1]['supplier'], 'Valve Co')        # from the task
        self.assertEqual(rows[0]['review'], 'approved')
        self.assertEqual(rows[2]['review'], 'pending')
        self.assertIn('DAQ', rows[1]['missing_evidence'])

    def test_filters(self):
        rows = app.inspection_report_rows()
        self.assertEqual(len(app.filter_report_rows(rows, {'supplier': 'Casting Ltd'})[0]), 2)
        self.assertEqual(len(app.filter_report_rows(rows, {'date_from': '2026-10-02', 'date_to': '2026-10-03'})[0]), 1)
        self.assertEqual(len(app.filter_report_rows(rows, {'inspector': 'Mr. Yu', 'result': 'Pass'})[0]), 1)
        self.assertEqual(len(app.filter_report_rows(rows, {'review': 'pending'})[0]), 2)

    def test_page_for_everyone(self):
        for name in ('hq', 'yu'):
            page = self.client_for(name).get('/reports?supplier=Casting+Ltd').get_data(as_text=True)
            self.assertIn('Casting Ltd', page)
            self.assertIn('/reports/export.xlsx?supplier=Casting+Ltd', page)

    def test_excel_export(self):
        response = self.client_for('hq').get('/reports/export.xlsx?supplier=Casting+Ltd')
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.get_data()))
        ws = wb['Inspections']
        self.assertEqual(ws.max_row, 3)                           # header + 2 reports
        headers = [c.value for c in ws[1]]
        notes = ws.cell(row=3, column=headers.index('备注 Notes') + 1).value
        self.assertTrue(notes.startswith("'="))                 # no formula injection
        summary = list(wb['Summary'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary, [('Casting Ltd', 2, 1, 1, 0, 50.0)])


if __name__ == '__main__':
    unittest.main()
