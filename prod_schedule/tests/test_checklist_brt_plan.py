"""BRT-plan checklists (L-Type hydrant heads): import layout, text answers,
item-code matching and replacing the paper checklist upload."""
import io
import json
import os
import sys
import tempfile
import unittest

os.environ.setdefault('APP_DATA_DIR', tempfile.mkdtemp(prefix='prod-schedule-tests-'))
os.environ.setdefault('SECRET_KEY', 'test-secret-key')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app
import checklists
from db import db_conn
from helpers import assign_job
from openpyxl import Workbook

JOB = 'MELBOURNE|PO-8|ACLTYPEDCFA'


def brt_workbook():
    wb = Workbook()
    ws = wb.active
    ws['B2'] = 'L-Type Hydrant - Batch Release Test Report'
    ws.append([])
    ws['B7'], ws['C7'], ws['E7'], ws['F7'], ws['G7'] = 'Test', 'Description', 'Checking Frequency per DPL', 'Result', 'Criteria'
    ws['K7'], ws['L7'], ws['M7'], ws['N7'] = 'Test 中文', 'Description 中文', 'Frequency 中文', 'Action 中文'
    rows = {
        9: ('Body Marking #1', 'Is "DAEMCO" marked on the body? [Y/N]', 'Single: 1pcs', 'If No, Reject & Check 100%',
            '本体标识 #1', '是否标有 DAEMCO？', '单口：1 件', '拒收并 100% 检查'),
        10: ('Body Marking #2', 'What is date of manufacture marked on the body? [MMYY]', 'Single: 1pcs', '-',
             '本体标识 #2', '生产日期？', '单口：1 件', ''),
        12: ('Checking point #1', 'Use plunger checking tool. Does it pass? [Y/N] - CRITICAL', 'Single & Dual: 100%',
             'If Yes, bolt needs to be tightened more', '检查点 #1', '工具能否通过？', '100%', '螺栓需要再拧紧'),
        13: ('Performance #8', 'DUAL ONLY: Can the end cap be screwed on? [Y/N]', 'Dual Only: 2pcs',
             'If fail, advise Murphy or Jayson.', '性能 #8', '仅双口：端盖能否拧上？', '仅双口：2 件', '通知 Murphy'),
        15: ('Pressure Testing #1', 'Test at 1800kPa; the video records 1 minute. Does the product leak? [Y/N]',
             'Single & Dual: 100%', 'If Yes, Reject', '压力测试 #1', '产品是否泄漏？', '100%', '拒收'),
    }
    for r, (label, text, freq, crit, label_zh, text_zh, freq_zh, action_zh) in rows.items():
        ws[f'B{r}'], ws[f'C{r}'], ws[f'E{r}'], ws[f'G{r}'] = label, text, freq, crit
        ws[f'K{r}'], ws[f'L{r}'], ws[f'M{r}'], ws[f'N{r}'] = label_zh, text_zh, freq_zh, action_zh
    ws['B17'] = 'Reliable Signature:'
    ws['B18'] = 'Should not be read'
    ws['C18'] = 'ignored'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class BrtPlanParseTests(unittest.TestCase):
    def test_layout(self):
        t = checklists.parse_checklist_workbook(brt_workbook())
        self.assertEqual(t['title'], 'L-Type Hydrant - Batch Release Test Report')
        self.assertEqual([(s['name'], s['name_zh'], len(s['questions'])) for s in t['sections']],
                         [('Body Marking', '本体标识', 2), ('Checking point', '检查点', 2), ('Pressure Testing', '压力测试', 1)])
        marking, date = t['sections'][0]['questions']
        self.assertEqual((marking['type'], marking['fail_on'], marking['action'], marking['hint']),
                         ('yes_no', 'no', 'Reject & Check 100%', 'Single: 1pcs'))
        self.assertEqual(marking['text'], 'Body Marking #1 — Is "DAEMCO" marked on the body?')
        self.assertEqual(marking['text_zh'], '本体标识 #1 — 是否标有 DAEMCO？')
        self.assertEqual(date['type'], 'text')
        tool, dual = t['sections'][1]['questions']
        self.assertEqual((tool['fail_on'], tool['action']), ('yes', 'bolt needs to be tightened more'))
        self.assertEqual((dual['fail_on'], dual['optional'], dual['action_zh']), ('no', True, '通知 Murphy'))
        self.assertEqual((dual['only_if'], tool['only_if']), ('Dual', ''))
        self.assertTrue(checklists.applies(dual, 'Dual CFA L - Type Hydrant Head'))
        self.assertFalse(checklists.applies(dual, 'Single CFA L - Type Hydrant Head'))
        self.assertEqual(t['sections'][2]['questions'][0]['photo'], 'always')      # the test video

    def test_text_answers(self):
        q = {'type': 'text'}
        self.assertEqual(checklists.evaluate(q, '0925'), 'ok')
        self.assertEqual(checklists.evaluate(q, ''), '')


class ItemCodeTemplateTests(unittest.TestCase):
    def setUp(self):
        app.app.config.update(TESTING=True)
        with db_conn() as conn:
            for table in ('users', 'inspection_tasks', 'checklist_templates', 'checklist_versions',
                          'inspection_drafts', 'draft_files', 'product_reference'):
                conn.execute(f'DELETE FROM {table}')
            conn.execute("INSERT INTO users (username,password_hash,role) VALUES ('murphy','x','admin')")
            conn.execute("INSERT INTO users (username,password_hash,role) VALUES ('yu','x','inspector')")
            self.ids = {r['username']: r['id'] for r in conn.execute('SELECT id, username FROM users')}
        schedule = {'MELBOURNE': [['Order Number', 'Daemco Purchase Order', 'Item Code', 'Item Description', 'Quantity'],
                                  ['DPL8', 'PO-8', 'ACLTYPEDCFA', 'Dual CFA L - Type Hydrant Head', '10'],
                                  ['DPL9', 'PO-9', 'ACLTYPE', 'L-type hydrant cover', '10'],
                                  ['DPL7', 'PO-7', 'ACLTYPESMFB', 'Single MFB L - Type Hydrant Head', '10']]}
        app.save_json(app.CURRENT_FILE, schedule)
        app.save_json(app.PREVIOUS_FILE, schedule)
        app.save_json(app.INSPECTIONS_CACHE, {})
        assign_job(JOB, 'yu')

    def client_for(self, name):
        client = app.app.test_client()
        with client.session_transaction() as session:
            session['user_id'] = self.ids[name]
            session['_csrf_token'] = 'tok'
        return client

    def test_item_codes_select_the_template_and_replace_the_paper_upload(self):
        murphy = self.client_for('murphy')
        murphy.post('/checklists/import', data={'_csrf_token': 'tok', 'target': 'new', 'name': 'L-Type heads',
                                                'file': (io.BytesIO(brt_workbook()), 'l.xlsx')},
                    content_type='multipart/form-data')
        with db_conn() as conn:
            tid = conn.execute('SELECT id FROM checklist_templates').fetchone()[0]
        _tpl, _ver, data = app.checklist_version(tid)
        murphy.post(f'/checklists/{tid}/edit', data={
            '_csrf_token': 'tok', 'name': 'L-Type heads', 'active': '1', 'data_json': json.dumps(data),
            'item_codes': 'ACLTYPESCFA, acltypesmfb\nACLTYPED*'})
        with db_conn() as conn:
            self.assertEqual(conn.execute('SELECT item_codes FROM checklist_templates').fetchone()[0],
                             'ACLTYPESCFA, ACLTYPESMFB, ACLTYPED*')
        self.assertEqual(app.checklist_for_item('ACLTYPEDCFA', '')['id'], tid)
        self.assertEqual(app.checklist_for_item('ACLTYPESMFB', '')['id'], tid)
        self.assertIsNone(app.checklist_for_item('ACLTYPE', 'L-type hydrant cover'))     # the cover is not a head

        page = self.client_for('yu').get(f'/inspect/{JOB}').get_data(as_text=True)
        self.assertIn('id="checklist-config"', page)
        self.assertNotIn('data-ev="checklist"', page)        # no paper checklist upload any more
        self.assertNotIn('该产品尚未配置所需证据', page)
        cover = self.client_for('murphy').get('/inspect/MELBOURNE|PO-9|ACLTYPE').get_data(as_text=True)
        self.assertIn('data-ev="checklist"', cover)            # the cover still uploads its checklist

        # submitting: no "missing checklist evidence"
        yu = self.client_for('yu')
        photo = yu.post(f'/inspect/{JOB}/draft/files', headers={'X-CSRF-Token': 'tok'}, data={
            'ref': 'product', 'template_id': tid, 'version': 1, 'file': (io.BytesIO(b'x'), 'p.jpg')},
            content_type='multipart/form-data')
        self.assertEqual(photo.status_code, 200)
        video_ref = data['sections'][2]['questions'][0]['id']
        yu.post(f'/inspect/{JOB}/draft/files', headers={'X-CSRF-Token': 'tok'}, data={
            'ref': video_ref, 'template_id': tid, 'version': 1, 'file': (io.BytesIO(b'x'), 'leak.mp4')},
            content_type='multipart/form-data')
        answers = {}
        for s in data['sections']:
            for q in s['questions']:
                answers[q['id']] = {'v': '0925' if q['type'] == 'text' else ('no' if q['fail_on'] == 'yes' else 'yes')}
        response = yu.post(f'/inspect/{JOB}/submit', data={
            '_csrf_token': 'tok', 'region': 'MELBOURNE', 'order_number': 'DPL8', 'item_code': 'ACLTYPEDCFA',
            'item_description': 'Dual CFA L - Type Hydrant Head', 'inspection_date': '2026-10-09',
            'quantity_inspected': '10', 'result': 'Pass', 'checklist_template_id': tid, 'checklist_version': 1,
            'checklist_json': json.dumps(answers)}, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 302)
        record = app.load_json(app.INSPECTIONS_CACHE, {})[JOB][0]
        self.assertEqual(record['missing_evidence'], [])
        self.assertEqual(record['checklist']['counts']['fail'], 0)

        # a checklist version that does not exist is refused, not ignored
        yu.post(f'/inspect/{JOB}/submit', data={
            '_csrf_token': 'tok', 'region': 'MELBOURNE', 'order_number': 'DPL8', 'item_code': 'ACLTYPEDCFA',
            'inspection_date': '2026-10-09', 'quantity_inspected': '10', 'result': 'Pass',
            'checklist_template_id': tid, 'checklist_version': 99, 'checklist_json': '{}'},
            content_type='multipart/form-data')
        self.assertEqual(len(app.load_json(app.INSPECTIONS_CACHE, {})[JOB]), 1)


class SplitLTypeTests(unittest.TestCase):
    def test_old_head_reports_get_the_head_type(self):
        app.save_json(app.INSPECTIONS_CACHE, {
            'M|P1|ACLTYPESCFA': [{'item_code': 'ACLTYPESCFA', 'item_description': 'Single CFA L - Type Hydrant Head',
                                  'product_type': 'l_type'}],
            'M|P2|ACLTYPE': [{'item_code': 'ACLTYPE', 'item_description': 'L-type hydrant cover',
                              'product_type': 'l_type'}]})
        self.assertEqual(app.reclassify_l_type_heads(), 1)
        cache = app.load_json(app.INSPECTIONS_CACHE, {})
        self.assertEqual(cache['M|P1|ACLTYPESCFA'][0]['product_type'], 'l_type_head')
        self.assertEqual(cache['M|P2|ACLTYPE'][0]['product_type'], 'l_type')
        self.assertEqual(app.reclassify_l_type_heads(), 0)
        with app.app.test_request_context():
            app.g.lang = 'zh'
            self.assertEqual(app.product_type_name('l_type_head'), 'L 型消防栓头')
            self.assertEqual(app.product_type_name('l_type'), 'L 型消防栓盖')


class SingleHeadTests(ItemCodeTemplateTests):
    SINGLE = 'MELBOURNE|PO-7|ACLTYPESMFB'
    test_item_codes_select_the_template_and_replace_the_paper_upload = None     # covered above

    def test_dual_only_questions_are_na_for_a_single_head_and_reports_upload(self):
        murphy = self.client_for('murphy')
        murphy.post('/checklists/import', data={'_csrf_token': 'tok', 'target': 'new', 'name': 'L-Type heads',
                                                'file': (io.BytesIO(brt_workbook()), 'l.xlsx')},
                    content_type='multipart/form-data')
        with db_conn() as conn:
            tid = conn.execute('SELECT id FROM checklist_templates').fetchone()[0]
        _tpl, _ver, data = app.checklist_version(tid)
        murphy.post(f'/checklists/{tid}/edit', data={'_csrf_token': 'tok', 'name': 'L', 'active': '1',
                                                    'data_json': json.dumps(data), 'item_codes': 'ACLTYPES*, ACLTYPED*'})
        page = murphy.get(f'/inspect/{self.SINGLE}').get_data(as_text=True)
        self.assertIn('"description": "Single MFB L - Type Hydrant Head"', page)

        def upload(ref, name):
            return murphy.post(f'/inspect/{self.SINGLE}/draft/files', headers={'X-CSRF-Token': 'tok'}, data={
                'ref': ref, 'template_id': tid, 'version': 1, 'file': (io.BytesIO(b'%PDF-1.4'), name)},
                content_type='multipart/form-data')
        self.assertEqual(upload('product', 'p.jpg').status_code, 200)
        video_ref = data['sections'][2]['questions'][0]['id']
        report = upload(video_ref, 'material-report.pdf')                 # documents are accepted now
        self.assertEqual(report.status_code, 200)
        self.assertTrue(report.get_json()['file']['doc'])
        self.assertEqual(upload(video_ref, 'tool.exe').status_code, 400)

        dual = data['sections'][1]['questions'][1]
        answers = {q['id']: {'v': '0925' if q['type'] == 'text' else ('no' if q['fail_on'] == 'yes' else 'yes')}
                   for s in data['sections'] for q in s['questions'] if q['id'] != dual['id']}
        answers[dual['id']] = {'v': 'no'}           # a stale "fail" for a question that does not apply
        response = murphy.post(f'/inspect/{self.SINGLE}/submit', data={
            '_csrf_token': 'tok', 'region': 'MELBOURNE', 'order_number': 'DPL7', 'item_code': 'ACLTYPESMFB',
            'item_description': 'Single MFB L - Type Hydrant Head', 'inspection_date': '2026-10-09',
            'quantity_inspected': '10', 'result': 'Pass', 'inspector_name': 'Murphy',
            'checklist_template_id': tid, 'checklist_version': 1, 'checklist_json': json.dumps(answers)},
            content_type='multipart/form-data')
        self.assertEqual(response.status_code, 302)
        record = app.load_json(app.INSPECTIONS_CACHE, {})[self.SINGLE][0]
        self.assertEqual(record['checklist']['answers'][dual['id']]['v'], 'na')
        self.assertEqual(record['checklist']['counts']['fail'], 0)


if __name__ == '__main__':
    unittest.main()
