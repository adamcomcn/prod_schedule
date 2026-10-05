"""Task list Excel export and the Monday summary e-mail."""
import io
import os
import sys
import tempfile
import unittest
from datetime import date, timedelta
from unittest import mock

os.environ.setdefault('APP_DATA_DIR', tempfile.mkdtemp(prefix='prod-schedule-tests-'))
os.environ.setdefault('SECRET_KEY', 'test-secret-key')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app
from db import db_conn
from openpyxl import load_workbook


class TasksExportTests(unittest.TestCase):
    def setUp(self):
        app.app.config.update(TESTING=True)
        today = app.china_today()
        with db_conn() as conn:
            for table in ('users', 'inspection_tasks', 'inspection_reviews'):
                conn.execute(f'DELETE FROM {table}')
            conn.execute("INSERT INTO users (username,password_hash,role,display_name) VALUES ('murphy','x','lead','Murphy')")
            conn.execute("INSERT INTO users (username,password_hash,role,display_name) VALUES ('yu','x','inspector','Mr. Yu')")
            self.ids = {r['username']: r['id'] for r in conn.execute('SELECT id, username FROM users')}
            tasks = [
                ('MEL|PO-1|A', 'D1', 'MEL', 'A', 'Casting Ltd', (today - timedelta(days=5)).isoformat(), 'Pending', self.ids['yu']),
                ('MEL|PO-2|B', 'D2', 'MEL', 'B', 'Valve Co', (today + timedelta(days=3)).isoformat(), 'In Progress', self.ids['yu']),
                ('QLD|PO-3|C', 'D3', 'QLD', 'C', '', (today + timedelta(days=40)).isoformat(), 'Pending', None),
                ('QLD|PO-4|D', 'D4', 'QLD', 'D', 'Valve Co', today.isoformat(), 'Completed', self.ids['yu']),
            ]
            for jk, order, region, item, sup, est, status, who in tasks:
                conn.execute('INSERT INTO inspection_tasks (job_key, order_number, region, item_code, supplier, '
                             'est_completion, status, assigned_to) VALUES (?,?,?,?,?,?,?,?)',
                             (jk, order, region, item, sup, est, status, who))
        app.save_json(app.INSPECTIONS_CACHE, {'QLD|PO-4|D': [{'result': 'Pass', 'inspection_date': '2026-10-01',
                                                              'order_number': 'D4', 'item_code': 'D'}]})
        app.save_json(app.CURRENT_FILE, {})
        config = app.load_config()
        config.update(task_notify_emails='murphy@example.com', weekly_summary=True)
        config.pop('weekly_summary_last', None)
        app.save_json(app.CONFIG_FILE, config)

    def client_for(self, name):
        client = app.app.test_client()
        with client.session_transaction() as session:
            session['user_id'] = self.ids[name]
            session['_csrf_token'] = 'tok'
        return client

    def export(self, name, query):
        response = self.client_for(name).get('/tasks/export.xlsx?' + query)
        self.assertEqual(response.status_code, 200)
        ws = load_workbook(io.BytesIO(response.get_data()))['Tasks']
        return [dict(zip([c.value for c in ws[1]], [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

    def test_export_follows_filters(self):
        rows = self.export('murphy', 'scope=all')
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0]['采购单号 PO'], 'PO-1')
        self.assertEqual(rows[0]['距今天数 Days left'], -5)
        done = next(r for r in rows if r['订单号 Order'] == 'D4')
        self.assertEqual(done['检验结果 Result'], 'Pass')
        self.assertIsNone(done['距今天数 Days left'])                 # finished: no countdown
        self.assertEqual(len(self.export('murphy', 'scope=all&status=Pending')), 2)
        self.assertEqual(len(self.export('murphy', 'scope=all&supplier=Valve+Co')), 2)
        self.assertEqual(len(self.export('murphy', 'scope=all&supplier=__none__')), 1)
        self.assertEqual(len(self.export('yu', 'scope=mine')), 3)

    def test_task_page_has_export_link(self):
        page = self.client_for('murphy').get('/tasks?scope=all&status=Pending').get_data(as_text=True)
        self.assertIn('/tasks/export.xlsx?scope=all&amp;status=Pending', page)

    def test_weekly_summary(self):
        sent = []
        fake = lambda *a: sent.append(a) or (True, 'sent')
        monday = app.china_today() - timedelta(days=app.china_today().weekday())
        tuesday = monday + timedelta(days=1)
        with mock.patch.object(app, '_smtp_send', side_effect=fake), app.app.test_request_context():
            with mock.patch.object(app, 'china_today', return_value=tuesday):
                self.assertFalse(app.send_weekly_summary()[0])        # only on Mondays
            with mock.patch.object(app, 'china_today', return_value=monday):
                self.assertTrue(app.send_weekly_summary()[0])
                self.assertFalse(app.send_weekly_summary()[0])        # once per Monday
        subject, body, recipients = sent[0][:3]
        self.assertEqual(recipients, ['murphy@example.com'])
        self.assertIn('每周汇总', subject)
        self.assertIn('Unassigned: 1', body)
        self.assertIn('Mr. Yu', body)

    def test_weekly_summary_button_and_switch(self):
        with mock.patch.object(app, '_smtp_send', return_value=(True, 'sent')) as send:
            self.client_for('murphy').post('/settings/weekly-summary', data={'_csrf_token': 'tok'})
            send.assert_not_called()                                   # admins only
        config = app.load_config()
        config['weekly_summary'] = False
        app.save_json(app.CONFIG_FILE, config)
        with app.app.test_request_context():
            self.assertFalse(app.send_weekly_summary()[0])


if __name__ == '__main__':
    unittest.main()
